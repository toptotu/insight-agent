import io
import os
from typing import Tuple

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}


def extract_text_from_upload(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    ext = os.path.splitext(filename.lower())[1]
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use PDF, DOCX, or PPTX.")
    if ext == ".pdf":
        return _extract_pdf_text(file_bytes), ext
    if ext == ".docx":
        return _extract_docx_text(file_bytes), ext
    if ext == ".pptx":
        return _extract_pptx_text(file_bytes), ext
    raise ValueError("Unsupported file type.")


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text:
            parts.append(text)
    return "\n".join(parts).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text:
            parts.append(paragraph.text)
    return "\n".join(parts).strip()


def _extract_pptx_text(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                parts.append(text)
    return "\n".join(parts).strip()
